import json
import re
import sys
from copy import deepcopy

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT


# =========================================================
# TEXT HELPERS
# =========================================================

def clean_text(value):
    if value is None:
        return ""

    return (
        str(value)
        .replace("\r", "")
        .replace("\u00A0", " ")
        .strip()
    )


def remove_markdown(text):
    text = clean_text(text)

    text = re.sub(
        r"^#{1,6}\s*",
        "",
        text
    )

    text = text.replace("**", "")
    text = text.replace("__", "")

    if (
        len(text) > 2
        and text.startswith("*")
        and text.endswith("*")
    ):
        text = text[1:-1]

    return text.strip()


def is_bullet(text):
    return bool(
        re.match(
            r"^\s*[-*•●▪◦]\s+",
            clean_text(text)
        )
    )


def strip_bullet(text):
    return re.sub(
        r"^\s*[-*•●▪◦]\s+",
        "",
        clean_text(text)
    ).strip()


def contains_email(text):
    return bool(
        re.search(
            r"[\w.+-]+@[\w.-]+\.\w+",
            text
        )
    )


def contains_phone(text):
    return bool(
        re.search(
            r"(?:\+?1[\s().-]*)?"
            r"\(?\d{3}\)?"
            r"[\s.-]*\d{3}"
            r"[\s.-]*\d{4}",
            text
        )
    )


# =========================================================
# SECTION HEADINGS
# =========================================================

COMMON_HEADINGS = [
    "PROFESSIONAL EXPERIENCE",
    "PROFESSIONAL SUMMARY",
    "TECHNICAL SKILLS",
    "WORK EXPERIENCE",
    "CERTIFICATIONS",
    "EXPERIENCE",
    "EDUCATION",
    "PROJECTS",
    "SUMMARY",
    "SKILLS"
]


def looks_like_heading(text):
    value = remove_markdown(text)

    if not value:
        return False

    upper = value.upper()

    if upper in COMMON_HEADINGS:
        return True

    # Allows future sections such as:
    # AWARDS
    # PUBLICATIONS
    # ACHIEVEMENTS
    # CERTIFICATES
    if (
        value == upper
        and len(value) <= 60
        and not contains_email(value)
        and not contains_phone(value)
        and not re.search(r"\b20\d{2}\b", value)
    ):
        return True

    return False


def split_inline_heading(text):
    """
    Handles:

    SUMMARY
    text

    and

    SUMMARY AI Engineer with...
    """

    value = remove_markdown(text)

    upper = value.upper()

    for heading in sorted(
        COMMON_HEADINGS,
        key=len,
        reverse=True
    ):
        if upper == heading:
            return heading, ""

        prefix = heading + " "

        if upper.startswith(prefix):
            return (
                heading,
                value[len(heading):].strip()
            )

    return None, value


# =========================================================
# EXPERIENCE HELPERS
# =========================================================

ROLE_WORDS = [
    "engineer",
    "developer",
    "administrator",
    "analyst",
    "scientist",
    "consultant",
    "architect",
    "specialist",
    "manager"
]


def looks_like_job_header(text):
    value = remove_markdown(text)

    lower = value.lower()

    return (
        any(
            word in lower
            for word in ROLE_WORDS
        )
        and
        (
            "," in value
            or "|" in value
        )
    )


def looks_like_date(text):
    value = remove_markdown(text)

    return bool(
        re.search(
            r"\b(?:19|20)\d{2}\b",
            value
        )
    )


# =========================================================
# EXTRACT
# =========================================================

def extract_lines(
    docx_path,
    output_json_path
):
    doc = Document(docx_path)

    lines = []

    for paragraph in doc.paragraphs:
        text = clean_text(
            paragraph.text
        )

        if text:
            lines.append(text)

    with open(
        output_json_path,
        "w",
        encoding="utf-8"
    ) as f:
        json.dump(
            {"lines": lines},
            f,
            ensure_ascii=False,
            indent=2
        )


# =========================================================
# TEMPLATE STYLE SNAPSHOTS
# =========================================================

def snapshot_paragraph(paragraph):
    if paragraph is None:
        return None

    ppr = None
    rpr = None

    if paragraph._p.pPr is not None:
        ppr = deepcopy(
            paragraph._p.pPr
        )

    if paragraph.runs:
        if paragraph.runs[0]._r.rPr is not None:
            rpr = deepcopy(
                paragraph.runs[0]._r.rPr
            )

    text = clean_text(
        paragraph.text
    )

    bullet_symbol = ""

    for symbol in [
        "",
        "•",
        "●",
        "▪",
        "◦"
    ]:
        if text.startswith(symbol):
            bullet_symbol = symbol
            break

    has_numbering = False

    if paragraph._p.pPr is not None:
        has_numbering = (
            paragraph._p.pPr.numPr
            is not None
        )

    return {
        "ppr": ppr,
        "rpr": rpr,
        "bullet_symbol": bullet_symbol,
        "has_numbering": has_numbering
    }


def visible_paragraphs(doc):
    return [
        p
        for p in doc.paragraphs
        if clean_text(p.text)
    ]


def find_exact(doc, wanted):
    wanted = wanted.upper()

    for paragraph in doc.paragraphs:
        if (
            clean_text(
                paragraph.text
            ).upper()
            == wanted
        ):
            return paragraph

    return None


def next_visible(doc, paragraph):
    if paragraph is None:
        return None

    paragraphs = list(
        doc.paragraphs
    )

    try:
        index = paragraphs.index(
            paragraph
        )
    except ValueError:
        return None

    for item in paragraphs[
        index + 1:
    ]:
        if clean_text(
            item.text
        ):
            return item

    return None


def find_first_bullet(doc):
    for paragraph in doc.paragraphs:

        text = clean_text(
            paragraph.text
        )

        if any(
            text.startswith(symbol)
            for symbol in [
                "",
                "•",
                "●",
                "▪",
                "◦"
            ]
        ):
            return paragraph

        if (
            paragraph._p.pPr is not None
            and
            paragraph._p.pPr.numPr
            is not None
        ):
            return paragraph

    return None


def find_experience_header(doc):
    inside = False

    for paragraph in doc.paragraphs:

        text = clean_text(
            paragraph.text
        )

        upper = text.upper()

        if upper in {
            "PROFESSIONAL EXPERIENCE",
            "WORK EXPERIENCE",
            "EXPERIENCE"
        }:
            inside = True
            continue

        if (
            inside
            and looks_like_heading(text)
        ):
            break

        if (
            inside
            and looks_like_job_header(text)
        ):
            return paragraph

    return None


def find_skill_line(doc):
    heading = (
        find_exact(
            doc,
            "TECHNICAL SKILLS"
        )
        or
        find_exact(
            doc,
            "SKILLS"
        )
    )

    return next_visible(
        doc,
        heading
    )


def collect_template_styles(doc):
    visible = visible_paragraphs(
        doc
    )

    name = (
        visible[0]
        if len(visible) > 0
        else None
    )

    title = (
        visible[1]
        if len(visible) > 1
        else name
    )

    contact = (
        visible[2]
        if len(visible) > 2
        else title
    )

    summary_heading = (
        find_exact(
            doc,
            "SUMMARY"
        )
        or
        find_exact(
            doc,
            "PROFESSIONAL SUMMARY"
        )
    )

    summary_text = (
        next_visible(
            doc,
            summary_heading
        )
    )

    heading = (
        summary_heading
        or
        find_exact(
            doc,
            "PROFESSIONAL EXPERIENCE"
        )
        or
        find_exact(
            doc,
            "EDUCATION"
        )
        or
        find_exact(
            doc,
            "TECHNICAL SKILLS"
        )
    )

    education_heading = (
        find_exact(
            doc,
            "EDUCATION"
        )
    )

    education = (
        next_visible(
            doc,
            education_heading
        )
    )

    return {
        "name":
            snapshot_paragraph(name),

        "title":
            snapshot_paragraph(title),

        "contact":
            snapshot_paragraph(contact),

        "heading":
            snapshot_paragraph(heading),

        "summary":
            snapshot_paragraph(
                summary_text
                or title
            ),

        "normal":
            snapshot_paragraph(
                summary_text
                or title
            ),

        "job_header":
            snapshot_paragraph(
                find_experience_header(
                    doc
                )
            ),

        "bullet":
            snapshot_paragraph(
                find_first_bullet(
                    doc
                )
            ),

        "education":
            snapshot_paragraph(
                education
                or summary_text
                or title
            ),

        "skill":
            snapshot_paragraph(
                find_skill_line(
                    doc
                )
                or summary_text
                or title
            )
    }


# =========================================================
# APPLY FORMATTING
# =========================================================

def apply_paragraph_snapshot(
    paragraph,
    snapshot
):
    if not snapshot:
        return

    if snapshot["ppr"] is None:
        return

    existing = paragraph._p.pPr

    if existing is not None:
        paragraph._p.remove(
            existing
        )

    paragraph._p.insert(
        0,
        deepcopy(
            snapshot["ppr"]
        )
    )


def apply_run_snapshot(
    run,
    snapshot
):
    if not snapshot:
        return

    rpr = snapshot.get(
        "rpr"
    )

    if rpr is None:
        return

    existing = run._r.rPr

    if existing is not None:
        run._r.remove(
            existing
        )

    run._r.insert(
        0,
        deepcopy(rpr)
    )

    run.font.highlight_color = None


# =========================================================
# CLEAR OLD RESUME CONTENT
# =========================================================

def clear_resume_body(doc):
    """
    Keeps sectPr, page size, margins,
    page borders, headers and footers.

    Removes all old resume body content.
    """

    body = doc._element.body

    for child in list(body):

        if child.tag.endswith(
            "}sectPr"
        ):
            continue

        body.remove(child)


# =========================================================
# CREATE NORMAL PARAGRAPH
# =========================================================

def add_paragraph(
    doc,
    snapshot,
    text="",
    bold=None,
    italic=None
):
    paragraph = doc.add_paragraph()

    apply_paragraph_snapshot(
        paragraph,
        snapshot
    )

    if text:
        run = paragraph.add_run(
            remove_markdown(text)
        )

        apply_run_snapshot(
            run,
            snapshot
        )

        if bold is not None:
            run.bold = bold

        if italic is not None:
            run.italic = italic

        run.font.highlight_color = None

    return paragraph


# =========================================================
# BULLET
# =========================================================

def add_bullet_paragraph(
    doc,
    snapshot,
    text
):
    paragraph = doc.add_paragraph()

    apply_paragraph_snapshot(
        paragraph,
        snapshot
    )

    content = remove_markdown(
        strip_bullet(text)
    )

    visible_symbol = ""

    if snapshot:
        visible_symbol = snapshot.get(
            "bullet_symbol",
            ""
        )

    # If original uses Word numbering,
    # pPr already contains numbering.
    if (
        snapshot
        and snapshot.get(
            "has_numbering"
        )
    ):
        final_text = content

    else:
        symbol = (
            visible_symbol
            or "•"
        )

        final_text = (
            f"{symbol} {content}"
        )

    run = paragraph.add_run(
        final_text
    )

    apply_run_snapshot(
        run,
        snapshot
    )

    run.bold = False
    run.font.highlight_color = None

    return paragraph


# =========================================================
# SKILL LINE
# =========================================================

def add_skill_paragraph(
    doc,
    snapshot,
    text
):
    paragraph = doc.add_paragraph()

    apply_paragraph_snapshot(
        paragraph,
        snapshot
    )

    content = remove_markdown(
        strip_bullet(text)
    )

    if ":" not in content:

        run = paragraph.add_run(
            content
        )

        apply_run_snapshot(
            run,
            snapshot
        )

        run.bold = False
        run.italic = False
        run.underline = False
        run.font.highlight_color = None

        return paragraph

    label, skills = content.split(
        ":",
        1
    )

    label_run = paragraph.add_run(
        f"{label.strip()}:"
    )

    apply_run_snapshot(
        label_run,
        snapshot
    )

    label_run.bold = True
    label_run.italic = False
    label_run.underline = False
    label_run.font.highlight_color = None


    skills_run = paragraph.add_run(
        f" {skills.strip()}"
    )

    apply_run_snapshot(
        skills_run,
        snapshot
    )

    skills_run.bold = False
    skills_run.italic = False
    skills_run.underline = False
    skills_run.font.highlight_color = None

    return paragraph


# =========================================================
# JOB HEADER
# =========================================================

def add_job_header(
    doc,
    snapshot,
    header,
    date=""
):
    paragraph = doc.add_paragraph()

    apply_paragraph_snapshot(
        paragraph,
        snapshot
    )

    header_run = paragraph.add_run(
        remove_markdown(
            header
        )
    )

    apply_run_snapshot(
        header_run,
        snapshot
    )


    if date:

        paragraph.add_run(
            "\t"
        )

        date_run = paragraph.add_run(
            remove_markdown(
                date
            )
        )

        apply_run_snapshot(
            date_run,
            snapshot
        )

    return paragraph


# =========================================================
# PARSE PASTED RESUME
# =========================================================

def parse_resume_text(text):
    raw_lines = (
        str(text or "")
        .replace("\r", "")
        .split("\n")
    )

    items = []

    for raw in raw_lines:

        raw = clean_text(raw)

        if not raw:
            continue

        heading, remainder = (
            split_inline_heading(
                raw
            )
        )

        if heading:

            items.append({
                "type": "heading",
                "text": heading,
                "raw": heading
            })

            if remainder:
                items.append({
                    "type": "normal",
                    "text": remainder,
                    "raw": remainder
                })

            continue

        if is_bullet(raw):
            item_type = "bullet"

        elif looks_like_date(raw):
            item_type = "date"

        else:
            item_type = "normal"

        items.append({
            "type": item_type,
            "text": remove_markdown(
                raw
            ),
            "raw": raw
        })

    return items


# =========================================================
# TITLE + CONTACT
# =========================================================

def split_title_contact(text):
    value = remove_markdown(
        text
    )

    phone = re.search(
        r"(?:\+?1[\s().-]*)?"
        r"\(?\d{3}\)?"
        r"[\s.-]*\d{3}"
        r"[\s.-]*\d{4}",
        value
    )

    if phone:

        position = phone.start()

        title = (
            value[:position]
            .strip(" |-")
            .strip()
        )

        contact = (
            value[position:]
            .strip()
        )

        return title, contact

    if contains_email(value):

        email = re.search(
            r"[\w.+-]+@[\w.-]+\.\w+",
            value
        )

        position = email.start()

        title = (
            value[:position]
            .strip(" |-")
            .strip()
        )

        contact = (
            value[position:]
            .strip()
        )

        return title, contact

    return value, ""


# =========================================================
# BUILD FINAL RESUME
# =========================================================

def build_resume(
    master_path,
    output_path,
    tailored_text
):
    doc = Document(
        master_path
    )

    styles = collect_template_styles(
        doc
    )

    items = parse_resume_text(
        tailored_text
    )

    if not items:
        raise Exception(
            "Tailored resume text is empty."
        )

    # Keep the uploaded DOCX itself so page
    # borders/margins/header/footer remain.
    # Remove only old resume body content.
    clear_resume_body(
        doc
    )


    # =====================================================
    # NAME
    # =====================================================

    add_paragraph(
        doc,
        styles["name"],
        items[0]["text"]
    )

    index = 1


    # =====================================================
    # TITLE / CONTACT BEFORE FIRST SECTION
    # =====================================================

    while (
        index < len(items)
        and
        items[index]["type"]
        != "heading"
    ):

        value = items[
            index
        ]["text"]

        if (
            contains_email(value)
            or contains_phone(value)
        ):

            title, contact = (
                split_title_contact(
                    value
                )
            )

            if title:

                add_paragraph(
                    doc,
                    styles["title"],
                    title
                )

            if contact:

                add_paragraph(
                    doc,
                    styles["contact"],
                    contact
                )

        else:

            add_paragraph(
                doc,
                styles["title"],
                value
            )

        index += 1


    # =====================================================
    # BODY
    # =====================================================

    current_section = ""


    while index < len(items):

        item = items[index]

        item_type = item[
            "type"
        ]

        text = item[
            "text"
        ]

        raw = item[
            "raw"
        ]


        # -------------------------------------------------
        # HEADING
        # -------------------------------------------------

        if item_type == "heading":

            current_section = (
                text.upper()
            )

            add_paragraph(
                doc,
                styles["heading"],
                text,
                bold=True
            )

            index += 1

            continue


        # -------------------------------------------------
        # EXPERIENCE
        # -------------------------------------------------

        if current_section in {
            "PROFESSIONAL EXPERIENCE",
            "WORK EXPERIENCE",
            "EXPERIENCE"
        }:

            if (
                item_type == "normal"
                and
                looks_like_job_header(
                    text
                )
            ):

                date = ""

                if (
                    index + 1
                    < len(items)
                    and
                    items[
                        index + 1
                    ]["type"]
                    == "date"
                ):

                    date = items[
                        index + 1
                    ]["text"]

                    index += 1

                add_job_header(
                    doc,
                    styles[
                        "job_header"
                    ],
                    text,
                    date
                )

                index += 1

                continue


            if item_type == "bullet":

                add_bullet_paragraph(
                    doc,
                    styles["bullet"],
                    raw
                )

                index += 1

                continue


            add_paragraph(
                doc,
                styles["normal"],
                text
            )

            index += 1

            continue


        # -------------------------------------------------
        # TECHNICAL SKILLS
        # -------------------------------------------------

        if current_section in {
            "TECHNICAL SKILLS",
            "SKILLS"
        }:

            add_skill_paragraph(
                doc,
                styles["skill"],
                raw
            )

            index += 1

            continue


        # -------------------------------------------------
        # EDUCATION
        # -------------------------------------------------

        if current_section == "EDUCATION":

            add_paragraph(
                doc,
                styles["education"],
                text
            )

            index += 1

            continue


        # -------------------------------------------------
        # PROJECTS / CERTIFICATIONS
        # -------------------------------------------------

        if current_section in {
            "PROJECTS",
            "CERTIFICATIONS"
        }:

            if item_type == "bullet":

                add_bullet_paragraph(
                    doc,
                    styles["bullet"],
                    raw
                )

            else:

                add_paragraph(
                    doc,
                    styles["normal"],
                    text
                )

            index += 1

            continue


        # -------------------------------------------------
        # SUMMARY / ANY OTHER SECTION
        # -------------------------------------------------

        if item_type == "bullet":

            add_bullet_paragraph(
                doc,
                styles["bullet"],
                raw
            )

        else:

            add_paragraph(
                doc,
                styles["summary"],
                text
            )

        index += 1


    doc.save(
        output_path
    )


# =========================================================
# WRITE
# =========================================================

def write_docx(input_json_path):
    print("DYNAMIC TEMPLATE WRITER RUNNING")

    with open(
        input_json_path,
        "r",
        encoding="utf-8"
    ) as f:
        payload = json.load(f)

    master_path = payload[
        "masterDocxPath"
    ]

    output_path = payload[
        "outputPath"
    ]

    tailored_text = payload.get(
        "tailoredResumeText",
        ""
    )

    # DEBUG
    print("====================================")
    print(
        "TAILORED TEXT CHARACTERS:",
        len(tailored_text)
    )
    print(
        "TAILORED TEXT NEWLINES:",
        tailored_text.count("\n")
    )

    test_items = parse_resume_text(
        tailored_text
    )

    print(
        "PARSED ITEMS:",
        len(test_items)
    )

    for i, item in enumerate(
        test_items[:15],
        1
    ):
        print(
            i,
            item["type"],
            repr(item["text"][:100])
        )

    print("====================================")

    if not clean_text(
        tailored_text
    ):
        raise Exception(
            "Tailored resume text is empty."
        )

    build_resume(
        master_path,
        output_path,
        tailored_text
    )


# =========================================================
# MAIN
# =========================================================

def main():

    if len(sys.argv) < 2:
        raise Exception(
            "Missing mode"
        )

    mode = sys.argv[1]

    if mode == "extract":

        if len(sys.argv) < 4:
            raise Exception(
                "Usage: docx_writer.py "
                "extract input.docx output.json"
            )

        extract_lines(
            sys.argv[2],
            sys.argv[3]
        )

        return


    if mode == "write":

        if len(sys.argv) < 3:
            raise Exception(
                "Usage: docx_writer.py "
                "write input.json"
            )

        write_docx(
            sys.argv[2]
        )

        return


    raise Exception(
        f"Unknown mode: {mode}"
    )


if __name__ == "__main__":
    main()